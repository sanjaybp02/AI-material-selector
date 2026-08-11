from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "AI_Material_Selector_Release_Note_and_Research_Paper.docx"

BLUE = "0B2545"
CYAN = "00A6C7"
LIGHT = "EAF4F7"
GRAY = "59636E"
DARK = "17212B"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_table_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(w)); grid.append(col)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in('w:tcW')
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_run(run, size=10.5, color=DARK, bold=False, italic=False, font='Aptos'):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:ascii'), font)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), font)
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold; run.italic = italic

def para(doc, text='', style=None, size=10.5, color=DARK, bold=False, italic=False, align=None, before=0, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if align is not None: p.alignment = align
    if text:
        r = p.add_run(text); set_run(r, size, color, bold, italic)
    return p

def rich_para(doc, parts, style=None, after=6):
    p = doc.add_paragraph(style=style); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.12
    for text, opts in parts:
        r = p.add_run(text); set_run(r, **opts)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); set_run(r, {1:16,2:13,3:11.5}[level], BLUE if level < 3 else CYAN, True)
    return p

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text); set_run(r, 10.2)
    return p

def numbered(doc, text):
    p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text); set_run(r, 10.2)
    return p

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.style = 'Table Grid'
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, BLUE); c.text = ''
        r = c.paragraphs[0].add_run(h); set_run(r, 9.4, 'FFFFFF', True)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(value)); set_run(r, 9.1)
            if len(t.rows) % 2 == 0: shade(cells[i], 'F6F9FA')
    if widths: set_table_widths(t, widths)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def callout(doc, label, text, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.style = 'Table Grid'
    c=t.cell(0,0); shade(c, fill); set_cell_margins(c, top=130, start=170, bottom=130, end=170); c.text=''
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(label + '  '); set_run(r, 10.2, BLUE, True)
    r=p.add_run(text); set_run(r, 10.2)
    set_table_widths(t,[9360]); doc.add_paragraph().paragraph_format.space_after=Pt(2)

def page_break(doc): doc.add_page_break()

def add_header_footer(section):
    header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r=header.add_run('AI Material Selector  |  Product and Technical Documentation'); set_run(r, 8.5, GRAY)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('AI Material Selector  •  Release documentation'); set_run(r, 8.5, GRAY)

def main():
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(0.82); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.85); sec.right_margin=Inches(0.85)
    sec.header_distance=Inches(0.35); sec.footer_distance=Inches(0.35); add_header_footer(sec)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(DARK)
    for name, size, color, before, after in [('Heading 1',16,BLUE,15,7),('Heading 2',13,BLUE,11,5),('Heading 3',11.5,CYAN,8,3)]:
        s=styles[name]; s.font.name='Aptos Display'; s._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos Display'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

    # Cover
    para(doc,'PRODUCT RELEASE NOTE + TECHNICAL RESEARCH PAPER',size=10,color=CYAN,bold=True,after=12)
    para(doc,'AI Material Selector',size=30,color=BLUE,bold=True,after=5)
    para(doc,'AI-assisted engineering material selection from requirements to recommendation, cost, sustainability, visualization, and export',size=15,color=GRAY,after=22)
    callout(doc,'DOCUMENT PURPOSE','A release-ready product note paired with a research-paper-style description of the system, implemented workflows, data path, AI orchestration, and validation boundaries.')
    table(doc,['Document field','Value'],[
        ('Project','AI Material Selector'),('Repository','github.com/sanjaybp02/AI-material-selector'),('Application type','Streamlit engineering intelligence dashboard'),('Primary AI service','Google Gemini via google-genai'),('Current release scope','Lite and Advanced material recommendation workflows'),('Prepared','22 July 2026')],[2100,7260])
    para(doc,'Prepared from the implementation in the project repository, including app.py, modules/, materials.csv, README.md, requirements.txt, settings.json, and history.db.',size=9,color=GRAY,italic=True,after=0)
    page_break(doc)

    # Contents / executive summary
    heading(doc,'Document map',1)
    for x in ['Part I - Product Release Note','Part II - Research Paper','Appendix A - End-to-end workflow reference','Appendix B - Implementation inventory','Appendix C - Known limitations and recommended next steps']:
        bullet(doc,x)
    heading(doc,'Executive summary',1)
    para(doc,'AI Material Selector is a Streamlit application that translates natural-language engineering requirements into material recommendations constrained by a local materials database. It combines deterministic filtering and unit conversion with Gemini-based reasoning, then adds cost, mass, embodied-carbon, comparison, history, PDF, and CAD-oriented outputs. The design is intentionally decision-support oriented: the application accelerates screening and trade-off analysis while requiring engineering verification before production use.')
    callout(doc,'CORE VALUE','The system narrows a broad material search into a traceable workflow: requirement -> candidate set -> AI ranking -> calculated part metrics -> visual review -> export and history.')
    table(doc,['Capability','What the release provides','Primary implementation'],[
        ('Fast screening','Single best material from a natural-language prompt','Lite mode; get_single_recommendation'),
        ('Engineering analysis','Top-three ranked candidates with pros, cons, reasoning, and confidence','Advanced mode; get_top3_recommendations'),
        ('Constraint control','Category, strength, temperature, density, fatigue, hardness, manufacturing, compliance, carbon filters','modules/filters.py'),
        ('Decision metrics','Mass, estimated raw-material cost, carbon footprint, confidence','modules/cost_engine.py + app.py'),
        ('Decision communication','PDF report, datasheet view, charts, comparison, chat follow-up','modules/pdf_report.py + charts.py + app.py'),
        ('Reuse and auditability','Templates and SQLite-backed search history with CSV/Excel export','templates.py + history.py'),
    ],[1800,4200,3360])
    page_break(doc)

    # Release note
    heading(doc,'Part I - Product Release Note',1)
    heading(doc,'1. Release overview',2)
    para(doc,'This release packages AI Material Selector as an engineering-focused decision-support dashboard. It is intended for early-stage design, feasibility checks, material substitution, cost-aware screening, and sustainability-aware comparison. It supports both quick answers and a deeper analytical workflow without requiring users to manually browse every material record.')
    heading(doc,'2. What is included',2)
    for item in [
        'Lite mode for a single direct recommendation with confidence, reasoning, estimated cost, mass, carbon footprint, PDF export, and datasheet view.',
        'Advanced mode for top-three ranking, physical constraint packs, editable candidate data for session-level what-if analysis, charts, comparisons, history, and follow-up chat.',
        'Natural-language prompts that can describe loading, operating environment, manufacturing process, budget, sustainability, and compliance needs.',
        'Industry presets for Aerospace, Medical, Heavy Machinery, and Sustainable screening.',
        'Metric and Imperial display with converted strength, temperature, modulus, density, thermal conductivity, fatigue, carbon, cost, volume, and mass units.',
        'Three cost sources: CSV database pricing, AI market estimation, and Live MetalPrice API with database fallback if live pricing is unavailable.',
        'Engineering PDF reports with reasoning, key properties, part specifications, cost, carbon footprint, and limitations.',
        'SQLite search history with editable rows and CSV/Excel export, plus reusable prompt templates and custom template creation.',
    ]: bullet(doc,item)
    heading(doc,'3. User-facing workflow',2)
    numbered(doc,'Configure the application in the sidebar: Gemini API key, model, mode, and unit system.')
    numbered(doc,'Define requirements using a prompt, optional replacement material, reusable template, part volume, and cost source.')
    numbered(doc,'In Advanced mode, narrow the database with category filters, industry packs, physical sliders, and compliance checks.')
    numbered(doc,'Click Find materials. The application sends only the current filtered candidate data and requirement query to Gemini.')
    numbered(doc,'Review the recommendation or ranked candidates. Inspect reasoning, confidence, cost, mass, carbon, pros, cons, charts, and datasheets.')
    numbered(doc,'Ask follow-up questions about trade-offs, manufacturing, or alternatives in Advanced mode.')
    numbered(doc,'Download a PDF engineering report, inspect the session history, export records, and use the generated CAD specimen workflow where exposed by the application.')
    heading(doc,'4. Release acceptance checklist',2)
    table(doc,['Area','Acceptance behavior'],[
        ('Startup','Application loads materials.csv, initializes SQLite history, and displays the engineering dashboard.'),
        ('Readiness','Find materials remains unavailable until a Gemini key, non-empty prompt, and non-empty candidate set exist.'),
        ('AI response','AI output is requested as strict JSON and matched back to an exact or fuzzy database material name.'),
        ('Cost','Part cost is derived from volume, density, and selected price source; live-price failure falls back to database pricing.'),
        ('Exports','Recommendation reports can be downloaded as PDF; history can be exported as CSV or Excel.'),
        ('Failure guidance','Empty filters can be explained by Gemini with a suggestion for which constraint to relax.'),
    ],[2000,7360])
    heading(doc,'5. Release caveats',2)
    callout(doc,'IMPORTANT','AI confidence is a model-generated score, not a statistical probability. Cost, carbon, and property values are screening estimates. Verify manufacturer datasheets, standards, certificates, process limits, safety factors, and current supplier quotes before design release or production.', 'FFF5D6')
    page_break(doc)

    # Research paper
    heading(doc,'Part II - Research Paper',1)
    heading(doc,'Abstract',2)
    para(doc,'Material selection is a multi-objective engineering activity in which performance, manufacturability, cost, sustainability, compliance, and operating environment must be considered together. AI Material Selector presents a hybrid decision-support architecture that combines deterministic data operations with large-language-model reasoning. A local tabular material database is normalized into Metric or Imperial display units, reduced by explicit engineering constraints, and passed as trusted context to Gemini for either single-candidate or top-three recommendation generation. The application then reattaches model outputs to database rows, calculates part-level mass, cost, and embodied carbon, and exposes the result through visual, conversational, historical, and export workflows. This paper describes the system architecture, algorithms, user workflows, reliability controls, and limitations. The system is best understood as an explainable screening assistant rather than an autonomous design authority.')
    heading(doc,'Keywords',2)
    para(doc,'materials selection; engineering decision support; retrieval-constrained generation; Streamlit; Google Gemini; multi-objective screening; embodied carbon; cost estimation')

    heading(doc,'1. Introduction',2)
    para(doc,'Material selection traditionally requires engineers to combine requirements, handbooks, supplier data, calculations, and experience. The friction is not only the number of material records; it is the translation between natural-language goals and measurable constraints. A request such as “lightweight, high-yield, machinable, and suitable for an elevated-temperature bracket” contains multiple objectives and trade-offs. AI Material Selector addresses this translation step by combining a structured property database with an AI reasoning layer while preserving deterministic calculations for units and part-level metrics.')
    heading(doc,'1.1 Research objective',3)
    para(doc,'The objective is to evaluate a practical hybrid workflow in which the language model interprets requirements and ranks candidates, while the application controls the candidate universe, property transformations, calculations, state, and exports. The design seeks to improve speed and usability without presenting generated reasoning as certified engineering analysis.')
    heading(doc,'1.2 Contributions',3)
    for item in ['A two-mode interface that supports both rapid screening and detailed comparison.', 'A retrieval-constrained prompt path: deterministic filtering occurs before AI ranking.', 'A unified part-cost and mass calculation path for Metric and Imperial display.', 'A sustainability-aware recommendation prompt using embodied carbon when relevant.', 'A traceability layer through reasoning, property views, search history, and downloadable reports.']: bullet(doc,item)

    heading(doc,'2. System architecture',2)
    table(doc,['Layer','Responsibilities','Key files'],[
        ('Presentation','Sidebar settings, guided tour, requirement input, filters, tabs, metrics, chat, export controls','app.py; modules/ui.py'),
        ('Application orchestration','Session state, readiness checks, recommendation processing, database row matching, result assembly','app.py'),
        ('AI services','Gemini calls, retry and fallback behavior, JSON prompt contracts, chat, empty-filter explanation','modules/ai_engine.py'),
        ('Data services','CSV loading, unit conversion, column mapping, currency and volume labels','modules/data_loader.py'),
        ('Decision calculations','Live metal price lookup, price-source selection, mass and raw-material cost','modules/cost_engine.py'),
        ('Presentation analytics','Radar chart, scatter plot, and property heatmap preparation','modules/charts.py'),
        ('Persistence and export','SQLite history, template storage, PDF generation, CAD STEP specimen generation','history.py; templates.py; pdf_report.py; cad_export.py'),
    ],[1500,4900,2960])
    heading(doc,'2.1 Data-flow model',3)
    para(doc,'The application loads the local material table once per Streamlit execution and keeps the raw DataFrame as the reference dataset. A converted display DataFrame is derived for the selected unit system. In Advanced mode, the display DataFrame is filtered and may be edited in a session-only data editor. The filtered table is serialized to text and embedded in the Gemini prompt. The model returns a material name and explanation; the name is matched back to the reference display DataFrame, after which deterministic cost and carbon calculations are applied.')
    callout(doc,'DESIGN PRINCIPLE','The model chooses among a controlled candidate set; it does not invent an unconstrained material catalog. Database matching is the gate between generated text and downstream calculations.')

    heading(doc,'3. Detailed workflows',2)
    heading(doc,'3.1 Workflow A - Configuration and readiness',3)
    numbered(doc,'Load environment values from .env and settings.json. The API key may come from the sidebar, saved settings, or GEMINI_API_KEY.')
    numbered(doc,'Load the materials CSV and initialize the search_history SQLite table if needed.')
    numbered(doc,'Select a Gemini model from the configured model list and choose Lite or Advanced mode.')
    numbered(doc,'Select Metric or Imperial units in Advanced mode; Lite uses the saved unit preference.')
    numbered(doc,'Persist settings including model, mode, unit system, part volume, cost source, and active filter values.')
    numbered(doc,'Enable analysis only when an API key exists, the prompt is non-empty, and the filtered candidate set is non-empty.')

    heading(doc,'3.2 Workflow B - Requirement capture and templating',3)
    para(doc,'The prompt is the natural-language problem statement. Users can select a quick-start template, search the template library, create a custom template, or describe an alternative to an existing material. When an existing material is selected, the application prefixes the request with an explicit replacement instruction. The volume is captured separately because cost and mass calculations require a numeric part volume.')
    table(doc,['Input','Role in workflow','Example'],[
        ('Project requirements','Semantic engineering objectives and context','Lightweight drone frame, high yield strength, machinable, up to 80 C'),
        ('Existing material','Frames replacement search','Find a better alternative to 6061 aluminum'),
        ('Part volume','Converts material properties into part mass and raw-material cost','50 cm3'),
        ('Cost source','Selects database, AI-estimated, or live price path','Use CSV Database Pricing'),
    ],[2100,4300,2960])

    heading(doc,'3.3 Workflow C - Deterministic constraint filtering',3)
    para(doc,'Advanced mode applies filters before AI ranking. The filter module supports category selection, industry constraint packs, essential temperature and yield thresholds, and advanced controls including density, modulus, fatigue, hardness, machinability, thermal conductivity, corrosion resistance, weldability, UV resistance, embodied carbon, bio-compatibility, and food-grade status. Active filters are summarized as chips and the candidate count is surfaced to the user.')
    table(doc,['Filter family','Operational effect'],[
        ('Industry packs','Pre-populate engineering constraints for Aerospace, Medical, Heavy Machinery, or Sustainable screening.'),
        ('Performance','Keep rows meeting minimum temperature, yield, modulus, fatigue, hardness, thermal, and process-property thresholds.'),
        ('Physical / economic','Apply maximum density and embodied-carbon constraints.'),
        ('Compliance','Require Bio-compatible and/or Food Grade values of Yes.'),
        ('What-if editing','Expose the filtered DataFrame for session-only edits before the AI call.'),
    ],[2600,6760])
    para(doc,'If the resulting set is empty, analysis stops. The user may ask the AI to explain why the combination is rare or physically difficult and which constraint may be worth relaxing.',size=10)

    heading(doc,'3.4 Workflow D - AI recommendation generation',3)
    para(doc,'The AI engine constructs a prompt containing the user query and the string representation of the current filtered database. The prompt positions Gemini as a mechanical and environmental engineering assistant and explicitly asks it to consider standard properties plus fatigue, hardness, corrosion or UV resistance, weldability, compliance, and embodied carbon. Lite mode requests one JSON object. Advanced mode requests a JSON array of up to three ranked objects with confidence, reasoning, pros, and cons. If AI market estimation is selected, the prompt additionally asks for EstimatedCostINR.')
    numbered(doc,'Create a Gemini client from the supplied API key.')
    numbered(doc,'Serialize the candidate DataFrame and construct a strict JSON output prompt.')
    numbered(doc,'Call the selected model through retry logic with up to three attempts per model and stable fallback models for transient 503 or 429 conditions.')
    numbered(doc,'Strip accidental Markdown code fences and parse the response as JSON.')
    numbered(doc,'Normalize a single object into a list when required and cap Advanced output at three results.')
    numbered(doc,'Match each returned MaterialName against the display DataFrame using exact match first, then case-insensitive substring matching.')
    numbered(doc,'Discard unmatched records rather than calculating cost or carbon from an unknown row.')
    callout(doc,'RELIABILITY CONTROL','The strict JSON contract and database re-match reduce formatting drift and prevent arbitrary model-generated material names from flowing directly into calculations. They do not eliminate hallucination, stale property data, or incorrect engineering interpretation.')

    heading(doc,'3.5 Workflow E - Cost, mass, and carbon calculation',3)
    para(doc,'The cost engine is deterministic after a material row has been matched. In Metric mode, mass is calculated from volume in cm3 and density in g/cm3, then converted to kg. Total cost is mass in kg multiplied by INR/kg. In Imperial mode, mass is calculated from volume in in3 and density in lb/in3; cost is calculated in USD/lb. The display conversion uses a default INR-to-USD rate of 85.0 when converting database values.')
    para(doc,'The calculation path can use the CSV row price, an AI-estimated INR/kg override, or a live MetalPrice API price. Live lookup maps common material names to supported symbols and falls back to database pricing when the API is unavailable. Embodied carbon is calculated as mass in kg multiplied by the row carbon factor in kg CO2/kg; Imperial presentation converts the displayed value to lb CO2.')
    table(doc,['Mode','Mass relationship','Cost relationship'],[
        ('Metric','mass_kg = volume_cm3 x density_g/cm3 / 1000','total = mass_kg x cost_INR_per_kg'),
        ('Imperial','mass_lb = volume_in3 x density_lb/in3','total = mass_lb x cost_USD_per_lb'),
    ],[1800,3780,3780])

    heading(doc,'3.6 Workflow F - Results, explanation, and comparison',3)
    para(doc,'Lite results show one recommendation, confidence, calculated cost, mass, carbon footprint, confidence bar, engineering reasoning, PDF report, and datasheet dialog. Advanced results show a ranked summary table and expandable candidate cards with confidence, cost, mass, carbon, reasoning, pros, cons, PDF report, and datasheet. The Charts tab provides radar, scatter, and heatmap views; Compare allows up to four selected materials. These views support human inspection of trade-offs instead of relying only on the rank order.')

    heading(doc,'3.7 Workflow G - Conversational follow-up',3)
    para(doc,'After Advanced analysis, the application seeds chat history with the original query and a concise summary of the top materials. The hidden chat context includes the current material database, selected cost source, and calculated costs. Follow-up questions are appended to the conversation and sent to Gemini with prior turns and this context. The intended use is to ask about processing, trade-offs, alternatives, or why a candidate ranked above another.')

    heading(doc,'3.8 Workflow H - Reporting and export',3)
    numbered(doc,'Build a PDF report from the selected material, confidence, reasoning, cost, mass, volume, currency, carbon metric, and properties.')
    numbered(doc,'Sanitize characters that are not safe for the FPDF Latin-1 output path.')
    numbered(doc,'Render a branded report containing part specifications, key properties, engineering reasoning, and assumptions and limitations.')
    numbered(doc,'Allow the user to download the report with a material-specific filename in Advanced mode.')
    numbered(doc,'Expose the material datasheet as a property table in a dialog.')
    numbered(doc,'Persist the search summary in SQLite and allow interactive editing plus CSV and Excel export.')
    numbered(doc,'Use the CAD export module to generate a STEP AP214 specimen block with the recommended material embedded in product metadata where that function is wired into the active UI.')

    heading(doc,'3.9 Workflow I - History and reuse',3)
    para(doc,'Each completed analysis logs timestamp, shortened query, top material, confidence, estimated cost, volume, unit system, and number of results. The History tab supports editing, adding, and deleting rows through Streamlit data_editor, saving the edited table back to SQLite, clearing all history, importing a validated CSV, and exporting CSV or Excel. Templates are stored separately in templates.json when users create custom prompts.')
    page_break(doc)

    heading(doc,'4. Research methodology and evaluation plan',2)
    para(doc,'The current repository implements a functional prototype and workflow demonstration rather than a controlled scientific benchmark. A rigorous evaluation should therefore measure both technical behavior and decision-support usefulness.')
    heading(doc,'4.1 Evaluation questions',3)
    for q in ['Does deterministic filtering reduce the candidate set without excluding valid solutions?', 'How often does Gemini return valid JSON and an exact or fuzzy-matchable material name?', 'How closely do cost and mass estimates agree with independent calculations?', 'Do users make faster or more consistent first-pass selections with the system?', 'How does the presence of embodied carbon and compliance filters change the recommended shortlist?']: bullet(doc,q)
    heading(doc,'4.2 Suggested test protocol',3)
    numbered(doc,'Create a labeled benchmark set of engineering requirements with an expert-approved shortlist and rationale.')
    numbered(doc,'Run each query in Lite and Advanced modes across several model selections and repeated trials.')
    numbered(doc,'Record candidate-set size, JSON validity, match success, top-one and top-three agreement, confidence, cost, carbon, and latency.')
    numbered(doc,'Compare outputs against expert decisions and independent spreadsheet calculations.')
    numbered(doc,'Conduct a user study measuring time-to-shortlist, perceived usefulness, explanation clarity, and calibration of confidence.')
    numbered(doc,'Report failure cases separately: no candidate, unmatched name, malformed JSON, transient service error, stale price, and unsupported material price.')
    heading(doc,'4.3 Metrics',3)
    table(doc,['Metric','Definition'],[
        ('JSON validity rate','Valid structured AI responses / total AI calls.'),
        ('Database match rate','Responses whose MaterialName maps to a database row / parsed responses.'),
        ('Top-k agreement','Whether expert-approved material appears in the first k results.'),
        ('Constraint preservation','Share of returned materials satisfying the active deterministic filters.'),
        ('Cost error','Absolute or percentage difference from an independent reference calculation or supplier quote.'),
        ('Carbon traceability','Whether displayed carbon can be reproduced from mass x row carbon factor.'),
        ('Human efficiency','Time and interaction count required to reach an approved shortlist.'),
    ],[2500,6860])

    heading(doc,'5. Reliability, safety, and governance',2)
    heading(doc,'5.1 Failure handling',3)
    for item in ['Transient 503/429 and demand-limit errors trigger retries with backoff and fallback models.', 'Non-transient errors are surfaced to the user with targeted guidance for unavailable models, missing keys, or general failures.', 'Empty candidate sets stop analysis and provide an optional explanation workflow.', 'Unmatched AI material names are rejected from downstream cost and carbon calculations.', 'Live price failures fall back to database pricing and display a warning.']: bullet(doc,item)
    heading(doc,'5.2 Engineering safety boundary',3)
    para(doc,'The tool does not perform complete structural analysis, fatigue-life certification, thermal simulation, corrosion qualification, regulatory certification, supplier qualification, or process validation. Its output is a screening recommendation. Production decisions require authoritative property data, manufacturing-specific design allowables, applicable standards, environmental and loading cases, safety factors, and qualified engineering review.')
    heading(doc,'5.3 Data and privacy considerations',3)
    para(doc,'The API key can be stored in settings.json by the current application behavior, and prompts plus filtered database context are sent to the configured Gemini service. Deployments should review secret management, access control, retention, and whether project requirements contain confidential or export-controlled information. A production deployment should prefer environment or secret-manager storage over writing keys to a local settings file.')

    heading(doc,'6. Discussion',2)
    para(doc,'The hybrid architecture is valuable because it assigns different responsibilities to deterministic code and generative AI. Filtering, unit conversion, matching, cost arithmetic, persistence, and export are explicit and inspectable. Natural-language interpretation, trade-off explanation, and conversational refinement are delegated to Gemini. This division improves usability while retaining an auditable path for many of the numerical outputs. The main research challenge is calibration: a fluent explanation and confidence score can appear more certain than the underlying evidence. Future versions should ground every recommendation in visible property deltas, add citations or source provenance for material records, and distinguish model confidence from rule-based feasibility.')
    heading(doc,'6.1 Current limitations',3)
    for item in ['The local materials.csv is the central knowledge source; its breadth, freshness, and property provenance determine recommendation quality.', 'AI confidence is not calibrated probability and may be sensitive to prompt wording and model choice.', 'The default currency conversion rate is static unless a deployment changes the implementation.', 'Live MetalPrice API support is name-mapping based and limited to supported symbols and market behavior.', 'Fuzzy matching by substring can select an unintended row when material names are ambiguous.', 'PDF output uses a Latin-1-safe path and may simplify or replace some Unicode characters.', 'The application is single-user/session oriented with local SQLite persistence rather than a multi-user transactional data service.', 'The CAD specimen generator creates a standard block representation and embeds metadata; it is not a geometry-aware design model or manufacturing-ready part definition.']: bullet(doc,item)

    heading(doc,'7. Conclusion',2)
    para(doc,'AI Material Selector demonstrates a practical path for applying generative AI to engineering material screening without making the language model the sole source of truth. Its strongest feature is the workflow composition: natural-language requirements are transformed into a constrained candidate set, ranked with AI reasoning, calculated into part-level metrics, and exposed through visual, conversational, historical, and export surfaces. With stronger source provenance, confidence calibration, evaluation datasets, and secure deployment controls, the prototype can evolve into a more rigorous engineering decision-support platform.')
    page_break(doc)

    heading(doc,'Appendix A - End-to-end workflow reference',1)
    table(doc,['Stage','Input','Processing','Output','Failure gate'],[
        ('1. Configure','API key, model, mode, units','Load settings and data; initialize history','Ready dashboard','Missing key or invalid settings'),
        ('2. Define','Prompt, template, replacement, volume, price source','Build final query and part context','Search-ready request','Empty prompt'),
        ('3. Constrain','Categories, packs, thresholds, compliance','Filter and optionally edit candidate rows','Filtered candidate table','Zero candidates'),
        ('4. Recommend','Filtered table + query','Gemini JSON call with retry/fallback','One or three AI candidates','API/service/JSON failure'),
        ('5. Match','AI MaterialName','Exact then substring match','Trusted material rows','Unmatched candidate'),
        ('6. Calculate','Row properties + volume + price source','Mass, cost, carbon arithmetic','Part metrics','Unsupported live price fallback'),
        ('7. Explain','Reasoning, properties, pros/cons','Render cards, datasheet, charts, chat','Human-reviewable decision set','No result state'),
        ('8. Export','Selected candidate and metrics','PDF, history, CSV/Excel, CAD specimen path','Portable artifacts','Encoding or export error'),
    ],[1400,1900,2600,1900,1560])
    heading(doc,'Appendix B - Implementation inventory',1)
    table(doc,['File','Role'],[
        ('app.py','Streamlit entry point, state orchestration, user workflow, AI integration, result rendering.'),
        ('modules/ai_engine.py','Gemini prompts, JSON parsing, retry/backoff/fallback, chat, filter-failure explanation.'),
        ('modules/data_loader.py','CSV loading, unit conversion, column mappings, unit/currency helpers.'),
        ('modules/filters.py','Industry packs, engineering sliders, compliance filters, candidate counts and chips.'),
        ('modules/cost_engine.py','Live price lookup and part-level mass/cost calculations.'),
        ('modules/charts.py','Radar, scatter, and heatmap visualization functions.'),
        ('modules/pdf_report.py','FPDF engineering report generation and Latin-1 sanitization.'),
        ('modules/history.py','SQLite history, data editor, import, CSV export, Excel export.'),
        ('modules/templates.py','Default, preset, and custom prompt template management.'),
        ('modules/cad_export.py','STEP AP214 specimen block generation with material metadata.'),
        ('modules/ui.py','Theme, hero, status, guided-tour, confidence, and supporting UI components.'),
        ('materials.csv','Structured local material property and pricing dataset.'),
        ('settings.json / history.db','Local settings and search-history persistence.'),
    ],[2500,6860])
    heading(doc,'Appendix C - Recommended next steps',1)
    for item in ['Add dataset provenance, source date, standards references, and material-grade identifiers to every row.', 'Replace static exchange-rate behavior with a configurable or timestamped market-data policy.', 'Add schema validation with a typed response model and explicit numeric bounds for confidence and estimated cost.', 'Add deterministic scoring baselines so AI rankings can be compared against weighted property rules.', 'Build a benchmark suite and regression tests for filters, units, cost, carbon, matching, and exports.', 'Move secrets to environment or platform secret storage and avoid persisting API keys in plain local settings.', 'Add authentication and a shared database if the application becomes multi-user.', 'Integrate the CAD specimen generator into the visible results UI only after clarifying its intended engineering use and geometry scope.', 'Add an explicit human-approval step before report export for regulated or safety-critical workflows.']: bullet(doc,item)

    doc.core_properties.title='AI Material Selector - Release Note and Research Paper'
    doc.core_properties.subject='Product release documentation, architecture, and workflow research paper'
    doc.core_properties.author='AI Material Selector project'
    doc.core_properties.keywords='materials selection, engineering AI, Streamlit, Gemini, research paper'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__': main()
